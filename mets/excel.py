"""Выгрузка найденных лотов в Excel."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = "#,##0"
PCT = "0%"
DATE = "DD.MM.YYYY HH:MM"

# (заголовок, ключ в словаре лота, формат, ширина)
COLUMNS = [
    ("Новый", "new_mark", None, 7),
    ("Лот", "number", None, 17),
    ("Название", "title", None, 45),
    ("Категория", "categories", None, 22),
    ("Регион", "region", None, 20),
    ("Форма торгов", "form_short", None, 14),
    ("Статус", "status", None, 22),
    ("Начало торгов", "trade_start", DATE, 16),
    ("Начальная цена, ₽", "start_price", MONEY, 14),
    ("Текущая цена, ₽", "price_now", MONEY, 14),
    ("Разница, ₽", "difference", MONEY, 14),
    ("Снижение сейчас", "discount_now", PCT, 10),
    ("Мин. цена, ₽", "price_min", MONEY, 13),
    ("Макс. снижение", "discount_max", PCT, 10),
    ("Период", "period", None, 9),
    ("След. снижение", "next_date", DATE, 16),
    ("След. цена, ₽", "next_price", MONEY, 13),
    ("Заявки до", "deadline", DATE, 16),
    ("Дней осталось", "days_left", "0.0", 9),
    ("Цена была, ₽", "prev_price", MONEY, 13),
    ("Задаток, ₽", "deposit", MONEY, 12),
    ("Заявок", "bids", "0", 7),
    ("Площадь, м²", "area", "#,##0.0", 10),
    ("Цена за м², ₽", "price_per_m2", MONEY, 11),
    ("Кадастровые номера", "cadastral", None, 22),
    ("Год выпуска", "year", None, 8),
    ("Обременения", "encumbrance", None, 25),
    ("Должник", "debtor", None, 25),
    ("Управляющий", "manager", None, 22),
    ("Телефон", "organizer_phone", None, 17),
    ("Email", "organizer_email", None, 22),
    ("Дело №", "case_number", None, 14),
    ("Описание", "description", None, 70),
]

# Добавляются после «Макс. снижение», если в поиске задан период «Цена в период»
WINDOW_COLUMNS = [
    ("Цена в начале периода, ₽", "win_price_start", MONEY, 14),
    ("Мин. цена в периоде, ₽", "win_price_min", MONEY, 14),
    ("Мин. цена действует с", "win_min_date", DATE, 16),
    ("Снижение в периоде", "win_discount", PCT, 10),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
WINDOW_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NEW_FILL = PatternFill("solid", fgColor="E2EFDA")
CHEAPER_FILL = PatternFill("solid", fgColor="FFF2CC")


def _form_short(lot):
    form = (lot.get("trade_form") or "").lower()
    if "публичн" in form:
        return "Публичка"
    if "аукцион" in form:
        return "Аукцион"
    if "конкурс" in form:
        return "Конкурс"
    return lot.get("trade_form") or ""


def _header(ws, titles):
    ws.append(titles)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"


def _columns(window):
    if not window:
        return COLUMNS
    i = [c[1] for c in COLUMNS].index("discount_max") + 1
    return COLUMNS[:i] + WINDOW_COLUMNS + COLUMNS[i:]


def _lots_sheet(ws, lots, window):
    columns = _columns(window)
    _header(ws, [c[0] for c in columns])
    for lot in lots:
        lot["form_short"] = _form_short(lot)
        cheaper = lot.get("prev_price") and lot.get("price_now") and lot["price_now"] < lot["prev_price"]
        lot["new_mark"] = "NEW" if lot.get("is_new") else ("↓" if cheaper else "")
        ws.append([lot.get(key) for _, key, _, _ in columns])
        row = ws.max_row
        for col, (_, key, fmt, _) in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=col)
            if fmt:
                cell.number_format = fmt
            if lot.get("is_new"):
                cell.fill = NEW_FILL
            elif cheaper and key in ("new_mark", "price_now", "prev_price"):
                cell.fill = CHEAPER_FILL
        link = ws.cell(row=row, column=2)
        link.hyperlink = lot["url"]
        link.font = Font(color="0563C1", underline="single")

    for col, (_, key, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if lots:
        last = ws.max_row
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
        # Чем больше снижение — тем зеленее
        for key in ("discount_now", "discount_max", "win_discount"):
            if key not in [c[1] for c in columns]:
                continue
            letter = get_column_letter([c[1] for c in columns].index(key) + 1)
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{last}",
                ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                               end_type="num", end_value=0.9, end_color="63BE7B"),
            )


def _period_state(r, now):
    if r["end"] and r["end"] <= now:
        return "прошёл"
    if r["start"] and r["start"] <= now:
        return "ТЕКУЩИЙ"
    return "будущий"


def _schedule_sheet(ws, lots, window_range):
    _header(ws, ["Лот", "Название", "№ периода", "Начало", "Окончание", "Цена, ₽", "Снижение от начальной", "Задаток, ₽", "Период"])
    now = datetime.now()
    w_from, w_to = window_range or (None, None)
    for lot in lots:
        for r in lot.get("schedule") or []:
            drop = 1 - r["price"] / lot["start_price"] if r["price"] and lot.get("start_price") else None
            state = _period_state(r, now)
            ws.append([lot["number"], lot["title"], r["n"], r["start"], r["end"], r["price"], drop, r["deposit"], state])
            row = ws.max_row
            for col, fmt in ((4, DATE), (5, DATE), (6, MONEY), (7, PCT), (8, MONEY)):
                ws.cell(row=row, column=col).number_format = fmt
            ws.cell(row=row, column=1).hyperlink = lot["url"]
            in_window = bool((w_from or w_to) and r["start"] and r["end"]
                             and (not w_to or r["start"] <= w_to) and (not w_from or r["end"] > w_from))
            for cell in ws[row]:
                if state == "ТЕКУЩИЙ":
                    cell.font = Font(bold=True)
                if in_window:
                    cell.fill = WINDOW_FILL
    for col, width in enumerate([17, 45, 9, 16, 16, 14, 11, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:I{ws.max_row}"


def _section(ws, title):
    if ws.max_row > 1:
        ws.append([])
    ws.append([title])
    cell = ws.cell(row=ws.max_row, column=1)
    cell.font = Font(bold=True, size=13, color="1F4E78")


def _about_sheet(ws, summary, filters, log):
    """Что искали и что делали: итоги, все фильтры словами, журнал действий."""
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    _section(ws, "Итоги")
    for key, value in summary:
        ws.append([key, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    _section(ws, "Фильтры поиска")
    for key, value in filters:
        ws.append([key, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    if log:
        _section(ws, "Журнал")
        for when, text in log:
            ws.append([when, text])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_excel(path, lots, summary, filters, log, window=False, window_range=None):
    wb = Workbook()
    _lots_sheet(wb.active, lots, window)
    wb.active.title = "Лоты"
    _schedule_sheet(wb.create_sheet("График цен"), lots, window_range)
    _about_sheet(wb.create_sheet("Что искали"), summary, filters, log)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
